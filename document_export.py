from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


CYAN = "54DAFF"
PURPLE = "B441FF"
BORDER = "D9D9D9"


def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _font(run, size=9, bold=False, color="000000"):
    run.font.name = "Yu Gothic"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        _shade(cell, CYAN)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                _font(run, bold=True)
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].text = "" if value is None else str(value)
            cells[index].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cells[index].paragraphs:
                for run in paragraph.runs:
                    _font(run)
    doc.add_paragraph()
    return table


def build_project_docx(project: dict) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    title = doc.add_paragraph(style="Title")
    run = title.add_run(f"{project['name']} 広告設計書")
    _font(run, size=22, bold=True)
    doc.add_paragraph("本書は、広告プロジェクトで確定したキャンペーン、キーワード、予算案、広告文をまとめたものです。")

    doc.add_heading("1 プロジェクト概要", level=1)
    _table(doc, ["項目", "内容"], [
        ["商品・サービス", project.get("product", "")],
        ["広告目的", project.get("objective", "")],
        ["LP", project.get("lp_url", "")],
        ["ブランド", project.get("brand", "")],
        ["強み", project.get("strengths", "")],
        ["共通除外キーワード", "、".join(project.get("shared_negative_keywords", []))],
    ])

    doc.add_heading("2 キャンペーン構成", level=1)
    _table(doc, ["キャンペーン", "広告グループ数", "日予算", "予算案"], [
        [c["name"], len(c.get("ad_groups", [])), c.get("daily_budget") or "未確定", c.get("budget_choice", "未確定")]
        for c in project.get("campaigns", [])
    ])

    for campaign_index, campaign in enumerate(project.get("campaigns", []), start=1):
        doc.add_page_break()
        doc.add_heading(f"3.{campaign_index} {campaign['name']}", level=1)
        _table(doc, ["配信条件", "設定"], [
            ["目的", campaign.get("objective", "")],
            ["LP", campaign.get("lp_url") or project.get("lp_url", "")],
            ["地域", campaign.get("locations", "")],
            ["言語", campaign.get("language", "")],
            ["ネットワーク", campaign.get("network", "")],
            ["ターゲット", "、".join(campaign.get("targets", []))],
            ["除外ターゲット", "、".join(campaign.get("excluded_targets", []))],
            ["固有除外キーワード", "、".join(campaign.get("negative_keywords", []))],
        ])
        for group_index, group in enumerate(campaign.get("ad_groups", []), start=1):
            doc.add_heading(f"広告グループ {group_index} {group['name']}", level=2)
            doc.add_paragraph(f"ターゲット: {group.get('target', '')}")
            doc.add_paragraph(f"検索意図: {group.get('intent', '')}")
            selected_keywords = [row for row in group.get("keywords", []) if row.get("selected", True)]
            _table(doc, ["キーワード", "コピー用", "マッチ", "検索数", "競合性", "低額帯", "高額帯", "判定"], [
                [
                    row.get("keyword", ""), row.get("copy_text", ""), row.get("match_type", ""),
                    row.get("monthly_searches", "未取得"), row.get("competition", "未取得"),
                    row.get("low_bid", ""), row.get("high_bid", ""), row.get("status", ""),
                ]
                for row in selected_keywords
            ])
            _table(doc, ["見出し", "通常文字数", "Google換算", "判定"], [
                [row.get("text", ""), row.get("normal_count", ""), row.get("google_count", ""), "OK" if row.get("google_count", 0) <= 30 else "超過"]
                for row in group.get("headlines", [])
            ])
            _table(doc, ["説明文", "通常文字数", "Google換算", "判定"], [
                [row.get("text", ""), row.get("normal_count", ""), row.get("google_count", ""), "OK" if row.get("google_count", 0) <= 90 else "超過"]
                for row in group.get("descriptions", [])
            ])
            doc.add_paragraph("CTA候補: " + "、".join(group.get("ctas", [])))

    output = BytesIO()
    doc.save(output)
    return output.getvalue()
